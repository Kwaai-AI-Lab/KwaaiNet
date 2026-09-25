#!/usr/bin/env python3
"""Build the anonymized AIAS+ 2026 manuscript (.docx) on the EasyChair Word template.

    python3 projects/kwaai-knowledge/papers/aias2026/build_manuscript_docx.py           # extended version
    python3 projects/kwaai-knowledge/papers/aias2026/build_manuscript_docx.py --short   # 4-page submission

--short builds DreamRAG-AIAS2026-short.md, which carries its own figures and captions inline, into
submission54-4page.docx.

Source of truth: DreamRAG-AIAS2026-manuscript.md (text) and the "Draft captions" section of
../../plans/DreamRAG-AIAS2026-manuscript-plan.md (figure captions). Figures: figures/fig*.pdf|png, produced by
figures/gen_paper_figures.py. Template: _easychair_template.docx, the EasyChair "Microsoft Word
Manuscript Template" (an input to this build; styles only, its content is discarded).

Output: submission54-manuscript.docx (or submission54-4page.docx with --short) next to this script,
tracked so the submission is in the repo. Neutral file name, author metadata blanked. Edit the .md,
never the .docx.

pandoc writes its own style names; the post-pass maps them onto the template's: Heading 1/2 ->
Section/Subsection (auto-numbered by the template, so the .md's manual numbers are stripped), body
paragraphs -> Normal, captions -> caption, code -> a Code block style (the template's Monospaced is a character style), tables -> Table Grid.
"""
from __future__ import annotations

import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
SRC = HERE / "DreamRAG-AIAS2026-manuscript.md"
PLAN = HERE.parents[1] / "plans/DreamRAG-AIAS2026-manuscript-plan.md"
TEMPLATE = HERE / "_easychair_template.docx"
FIGS = HERE / "figures"
OUT = HERE / "submission54-manuscript.docx"
SHORT = False

# Figure n goes after the first paragraph of the manuscript containing this anchor.
FIGURE_ANCHORS = {
    1: "Figure 1 gives an overview.",
    2: "Figure 2 reports the three measures",
    3: "**Fact density.** Figure 3",
    4: "**Fact density.** Figure 3",
    5: "**Accuracy over cycles.**",
    6: "**Consolidation** is bounded",
    7: "Figure 7 shows answer recall",
}
FIGURE_FILES = {
    1: "fig1_architecture.png", 2: "fig2_cross_corpus.png", 3: "fig3_fact_density.png",
    4: "fig4_completeness_per_cycle.png", 5: "fig5_completeness_vs_accuracy.png",
    6: "fig6_cost.png", 7: "fig7_development_history.png",
}
FIGURE_WIDTH = {5: "60%"}  # the two-panel stacked figure is narrow by design


def captions() -> dict[int, str]:
    text = PLAN.read_text()
    block = text[text.index("### Draft captions"):]
    out = {}
    for m in re.finditer(r"^- \*\*Figure (\d)\.\*\* (.*?)(?=^- \*\*|^\S)", block, re.S | re.M):
        out[int(m[1])] = " ".join(m[2].split())
    missing = set(FIGURE_FILES) - set(out)
    if missing:
        raise SystemExit(f"no caption in the plan for figure(s) {sorted(missing)}")
    return out


def preprocess(md: str) -> str:
    md = re.sub(r"\A---\n.*?\n---\n", "", md, flags=re.S)       # YAML front matter
    md = re.sub(r"<!--.*?-->\n?", "", md, flags=re.S)             # internal notes
    title = re.search(r'^title: "(.*)"$', SRC.read_text(), re.M)[1]

    abstract = re.search(r"^## Abstract\n\n(.*?)(?=^## )", md, re.S | re.M)[1].strip()
    md = md[:md.index("## Abstract")] + md[md.index("## 1 Introduction"):]
    abs_paras, kw = [], ""
    for para in abstract.split("\n\n"):
        (kw := para) if para.startswith("**Keywords:**") else abs_paras.append(para)

    def div(style: str, body: str) -> str:
        return f'::: {{custom-style="{style}"}}\n{body}\n:::\n'

    front = [
        f"---\ntitle: \"{title}\"\n---\n",
        div("Authors", "Anonymous author(s)"),
        div("Institute", "Anonymous institution(s)"),
        div("Abstract title", "Abstract"),
        *[div("Abstract", p) for p in abs_paras],
        div("Abstract", kw),
    ]

    # Heading levels: "## 1 Intro" -> Section, "### 3.1 Ingestion" -> Subsection; strip numbers.
    md = re.sub(r"^## \d+ (.*)$", r"# \1", md, flags=re.M)
    md = re.sub(r"^### \d+\.\d+ (.*)$", r"## \1", md, flags=re.M)
    # Unnumbered back-matter headings use the template's References style.
    md = re.sub(r"^## (References|Appendix.*)$",
                lambda m: div("References", m[1]), md, flags=re.M)
    # Table captions: "**Table N.** ..." paragraphs -> caption style.
    md = re.sub(r"^(\*\*Table \d\.\*\*.*?)(?=\n\n)", lambda m: div("caption", m[1]), md,
                flags=re.S | re.M)
    # Reference list entries -> Bibliography style.
    ref_start = md.index('custom-style="References"}\nReferences')
    head, refs = md[:ref_start], md[ref_start:]
    refs = re.sub(r"^- (.*?)(?=\n- |\n*\Z)",
                  lambda m: div("Bibliography", " ".join(m[1].split())), refs, flags=re.S | re.M)
    md = head + refs

    if SHORT:  # the short text places its own figures
        return "\n".join(front) + "\n" + md
    # Figures, each after its anchor paragraph.
    caps = captions()
    paras = md.split("\n\n")
    for n in sorted(FIGURE_FILES, reverse=True):
        idx = next(i for i, p in enumerate(paras) if FIGURE_ANCHORS[n] in p)
        width = FIGURE_WIDTH.get(n, "100%")
        fig = f"![**Figure {n}.** {caps[n]}]({FIGS / FIGURE_FILES[n]}){{width={width}}}"
        paras.insert(idx + 1, fig)
    # Two figures share an anchor: keep them in numeric order.
    return "\n".join(front) + "\n" + "\n\n".join(paras)


# pandoc style IDs -> template style names. Matched on the raw w:pStyle id: pandoc's IDs are not
# defined in the template, so python-docx reports them all as "Normal".
STYLE_MAP = {
    "Heading1": "Section", "Heading2": "Subsection",
    "FirstParagraph": "Normal", "BodyText": "Normal",
    "ImageCaption": "Caption", "TableCaption": "Caption",
    "Compact": "List Paragraph",
}


def style_id(p) -> str | None:
    ppr = p._p.pPr
    return ppr.pStyle.val if ppr is not None and ppr.pStyle is not None else None


def ensure_style(doc, name: str, *, size: float, mono: bool = False):
    try:
        return doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        st.font.size = Pt(size)
        if mono:
            st.font.name = "Courier New"
        pf = st.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_after = Pt(0)
        return st


# Column widths (inches) for tables whose column count matches; the text area is 5.7 in wide.
COL_WIDTHS = {4: [2.5, 1.3, 1.35, 0.55]}


def set_col_widths(table, widths_in) -> None:
    """Fixed layout, with widths written to both tblGrid and every cell, so Word honours them."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    for old in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    # Schema order: tblLayout precedes tblCellMar, tblLook, tblCaption, tblDescription. Appending it
    # after them makes Word report "unreadable content".
    later = [tblPr.find(qn(f"w:{t}")) for t in ("tblCellMar", "tblLook", "tblCaption", "tblDescription")]
    later = [e for e in later if e is not None]
    if later:
        later[0].addprevious(layout)
    else:
        tblPr.append(layout)
    for col, w in zip(tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")), widths_in):
        col.set(qn("w:w"), str(int(w * 1440)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in):
            cell.width = Inches(w)


def postprocess(path: Path) -> None:
    doc = Document(path)
    mono = ensure_style(doc, "Code block", size=8, mono=True)
    table_text = ensure_style(doc, "Table text", size=8)
    for p in doc.paragraphs:
        sid = style_id(p)
        if sid in STYLE_MAP:
            p.style = doc.styles[STYLE_MAP[sid]]
        elif sid == "SourceCode":
            p.style = mono
        elif sid in ("CaptionedFigure", "Figure"):
            p.style = doc.styles["Normal"]
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for t in doc.tables:
        t.style = doc.styles["Table Grid"]
        if len(t.columns) in COL_WIDTHS:
            set_col_widths(t, COL_WIDTHS[len(t.columns)])
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    p.style = table_text
    cp = doc.core_properties
    cp.author = cp.last_modified_by = cp.comments = cp.keywords = cp.category = ""
    cp.title = "Submission 54"
    doc.save(path)


def main() -> None:
    global SRC, OUT, SHORT
    SHORT = "--short" in sys.argv[1:]
    if SHORT:
        SRC = HERE / "DreamRAG-AIAS2026-short.md"
        OUT = OUT.with_name("submission54-4page.docx")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        mid = Path(tmp) / "manuscript.md"
        mid.write_text(preprocess(SRC.read_text()))
        subprocess.run(["pandoc", str(mid), "-f", "markdown", "--reference-doc", str(TEMPLATE),
                        "--resource-path", str(HERE), "-o", str(OUT)], check=True)
    postprocess(OUT)
    print("wrote", OUT.relative_to(REPO))


if __name__ == "__main__":
    main()
